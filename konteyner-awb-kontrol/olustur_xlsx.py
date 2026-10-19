#!/usr/bin/env python3
"""konteyner-awb-kontrol.xlsx dosyasını üretir (openpyxl).

Kullanım:  python3 olustur_xlsx.py
Çıktı:     bu betiğin yanında konteyner-awb-kontrol.xlsx

Çalışma kitabı yalnızca formül içerir, makro yoktur. Formüller Excel 2010 ve LibreOffice
ile uyumludur (LET, XLOOKUP, FILTER, TEXTJOIN, INDIRECT, OFFSET kullanılmaz).

Lisans: MIT. LogisticsIQ hazırladı: https://logisticsiq.org
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

SURUM_TARIHI = "2026-10-19"
SITE = "https://logisticsiq.org"
ILK_SATIR = 5
SATIR_SAYISI = 1000
SON_SATIR = ILK_SATIR + SATIR_SAYISI - 1

HARFLER = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
RAKAMLAR = "0123456789"
# ISO 6346 harf değerleri: A=10, 11'in katları atlanır (A..Z için 26 değer).
HARF_DEGERLERI = "{10,12,13,14,15,16,17,18,19,20,21,23,24,25,26,27,28,29,30,31,32,34,35,36,37,38}"
AGIRLIKLAR = "{1,2,4,8,16,32,64,128,256,512}"

SARI = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
BASLIK_DOLGU = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
YESIL_DOLGU = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
KIRMIZI_DOLGU = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
KALIN = Font(name="Calibri", size=11, bold=True)
GRI = Font(name="Calibri", size=11, color="808080", italic=True)
BASLIK_FONT = Font(name="Calibri", size=14, bold=True)
BAGLANTI_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")


def temizle(hucre):
    """Boşluk, tire, nokta ve eğik çizgiyi siler."""
    ifade = hucre
    for karakter in (" ", "-", ".", "/"):
        ifade = 'SUBSTITUTE(%s,"%s","")' % (ifade, karakter)
    return ifade


def konteyner_formulleri(r):
    a, b, f = "$A%d" % r, "$B%d" % r, "$F%d" % r
    # Büyük harf: önce küçük i büyük I yapılır, çünkü Türkçe Excel UPPER("i") için "İ" verir.
    # Noktasız küçük ı geçersiz sayılır (Python ve VBA ile aynı davranış).
    norm = 'IF(%s="","",UPPER(SUBSTITUTE(SUBSTITUTE(%s,"i","I"),"ı","?")))' % (a, temizle(a))
    bicim = (
        'IF({a}="","",IF(AND(LEN({b})=11,'
        'SUMPRODUCT(--ISNUMBER(FIND(MID({b},{{1,2,3}},1),"{harf}")))=3,'
        'ISNUMBER(FIND(MID({b},4,1),"UJZ")),'
        'SUMPRODUCT(--ISNUMBER(FIND(MID({b},{{5,6,7,8,9,10,11}},1),"{rak}")))=7),1,0))'
    ).format(a=a, b=b, harf=HARFLER, rak=RAKAMLAR)
    deger = {}
    for sira in range(1, 5):  # 1-4. karakterler harf
        deger[sira] = 'IF({f}=1,INDEX({dizi},1,FIND(MID({b},{s},1),"{harf}")),"")'.format(
            f=f, dizi=HARF_DEGERLERI, b=b, s=sira, harf=HARFLER)
    for sira in range(5, 11):  # 5-10. karakterler rakam
        deger[sira] = 'IF({f}=1,FIND(MID({b},{s},1),"{rak}")-1,"")'.format(f=f, b=b, s=sira, rak=RAKAMLAR)
    toplam = 'IF({f}=1,SUMPRODUCT($G{r}:$P{r},{w}),"")'.format(f=f, r=r, w=AGIRLIKLAR)
    beklenen = 'IF({f}=1,MOD(MOD($Q{r},11),10),"")'.format(f=f, r=r)
    hesaplanan = 'IF({f}=1,$R{r},"")'.format(f=f, r=r)
    sonuc = (
        'IF({a}="","",IF({f}=0,"Biçim hatalı",IF(FIND(MID({b},11,1),"{rak}")-1=$R{r},"Geçerli",'
        '"Kontrol hanesi hatalı (doğrusu: "&$R{r}&")")))'
    ).format(a=a, f=f, b=b, rak=RAKAMLAR, r=r)
    return norm, hesaplanan, sonuc, bicim, deger, toplam, beklenen


def awb_formulleri(r):
    a, b, f = "$A%d" % r, "$B%d" % r, "$F%d" % r
    norm = 'IF(%s="","",%s)' % (a, temizle(a))
    bicim = (
        'IF({a}="","",IF(AND(LEN({b})=11,'
        'SUMPRODUCT(--ISNUMBER(FIND(MID({b},{{1,2,3,4,5,6,7,8,9,10,11}},1),"{rak}")))=11),1,0))'
    ).format(a=a, b=b, rak=RAKAMLAR)
    seri = 'IF({f}=1,VALUE(MID({b},4,7)),"")'.format(f=f, b=b)
    beklenen = 'IF({f}=1,MOD($G{r},7),"")'.format(f=f, r=r)
    hesaplanan = 'IF({f}=1,$H{r},"")'.format(f=f, r=r)
    sonuc = (
        'IF({a}="","",IF({f}=0,"Biçim hatalı",IF(FIND(MID({b},11,1),"{rak}")-1=$H{r},"Geçerli",'
        '"Kontrol hanesi hatalı (doğrusu: "&$H{r}&")")))'
    ).format(a=a, f=f, b=b, rak=RAKAMLAR, r=r)
    return norm, hesaplanan, sonuc, bicim, seri, beklenen


def sayfa_duzeni(ws, baslik, aciklama1, aciklama2, sutun_basliklari, yardimci_basliklar, ilk_yardimci, son_yardimci):
    ws["A1"] = baslik
    ws["A1"].font = BASLIK_FONT
    ws["A2"] = aciklama1
    ws["A3"] = aciklama2
    for satir in (2, 3):
        ws.merge_cells("A%d:D%d" % (satir, satir))
        ws["A%d" % satir].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[satir].height = 48
    for sutun, metin in zip("ABCD", sutun_basliklari):
        hucre = ws["%s4" % sutun]
        hucre.value = metin
        hucre.font = KALIN
        hucre.fill = BASLIK_DOLGU
        hucre.alignment = Alignment(wrap_text=True, vertical="center")
    ws["A4"].fill = SARI
    ws.row_dimensions[4].height = 32
    for sutun, metin in yardimci_basliklar.items():
        ws["%s4" % sutun] = metin
        ws["%s4" % sutun].font = GRI
    for sutun, genislik in zip("ABCD", (30, 24, 18, 44)):
        ws.column_dimensions[sutun].width = genislik
    ws.freeze_panes = "A5"
    # Yardımcı sütunları gruplayıp gizle (soldaki + düğmesiyle açılır).
    ws.column_dimensions.group(ilk_yardimci, son_yardimci, hidden=True)
    ws.sheet_properties.outlinePr.summaryRight = False


def sayfa_bicimle(ws):
    yesil = FormulaRule(formula=['$D%d="Geçerli"' % ILK_SATIR], fill=YESIL_DOLGU,
                        font=Font(color="006100"), stopIfTrue=False)
    kirmizi = FormulaRule(formula=['AND($D%d<>"",$D%d<>"Geçerli")' % (ILK_SATIR, ILK_SATIR)],
                          fill=KIRMIZI_DOLGU, font=Font(color="9C0006"), stopIfTrue=False)
    aralik = "A%d:D%d" % (ILK_SATIR, SON_SATIR)
    ws.conditional_formatting.add(aralik, yesil)
    ws.conditional_formatting.add(aralik, kirmizi)


def konteyner_sayfasi(wb):
    ws = wb.active
    ws.title = "Konteyner"
    sayfa_duzeni(
        ws,
        "Konteyner numarası kontrolü (ISO 6346)",
        "Konteyner numaralarını A5 hücresinden başlayarak A sütununa yapıştırın (en çok %d satır). "
        "Büyük ya da küçük harf, boşluk, tire, nokta ve eğik çizgi fark etmez." % SATIR_SAYISI,
        "Geçerli bir kontrol hanesi yalnızca numaranın biçimce doğru olduğunu gösterir. "
        "Konteynerin var olduğunu ya da o sahibe ait olduğunu göstermez.",
        ["Numara (sarı hücrelere yapıştırın)", "Temizlenmiş numara", "Hesaplanan kontrol hanesi", "Sonuç"],
        {"F": "Yardımcı: biçim doğru (1/0)",
         **{chr(ord("G") + i): "Yardımcı: %d. karakter değeri" % (i + 1) for i in range(10)},
         "Q": "Yardımcı: ağırlıklı toplam", "R": "Yardımcı: beklenen hane"},
        "F", "R")
    for r in range(ILK_SATIR, SON_SATIR + 1):
        norm, hesaplanan, sonuc, bicim, deger, toplam, beklenen = konteyner_formulleri(r)
        ws["A%d" % r].fill = SARI
        ws["A%d" % r].number_format = "@"
        ws["B%d" % r] = "=" + norm
        ws["C%d" % r] = "=" + hesaplanan
        ws["C%d" % r].number_format = "0"
        ws["C%d" % r].alignment = Alignment(horizontal="center")
        ws["D%d" % r] = "=" + sonuc
        for sutun in "BCD":
            ws["%s%d" % (sutun, r)].font = KALIN
        ws["F%d" % r] = "=" + bicim
        for sira in range(1, 11):
            ws["%s%d" % (chr(ord("G") + sira - 1), r)] = "=" + deger[sira]
        ws["Q%d" % r] = "=" + toplam
        ws["R%d" % r] = "=" + beklenen
    sayfa_bicimle(ws)


def awb_sayfasi(wb):
    ws = wb.create_sheet("AWB")
    sayfa_duzeni(
        ws,
        "Hava konşimentosu (AWB) numarası kontrolü",
        "AWB numaralarını A5 hücresinden başlayarak A sütununa yapıştırın (en çok %d satır). "
        "Numara 11 rakamdır: 3 haneli havayolu öneki ve 8 haneli seri numarası. "
        "Tire ve boşluk fark etmez (örnek: 020-12345675)." % SATIR_SAYISI,
        "Seri numarasının 8. hanesi, ilk 7 hanenin 7'ye bölümünden kalana eşit olmalıdır. "
        "Sayı olarak yapıştırırsanız baştaki sıfır kaybolabilir, bu yüzden sütun metin biçimindedir.",
        ["Numara (sarı hücrelere yapıştırın)", "Temizlenmiş numara", "Hesaplanan kontrol hanesi", "Sonuç"],
        {"F": "Yardımcı: biçim doğru (1/0)", "G": "Yardımcı: seri numarasının ilk 7 hanesi",
         "H": "Yardımcı: beklenen hane"},
        "F", "H")
    for r in range(ILK_SATIR, SON_SATIR + 1):
        norm, hesaplanan, sonuc, bicim, seri, beklenen = awb_formulleri(r)
        ws["A%d" % r].fill = SARI
        ws["A%d" % r].number_format = "@"
        ws["B%d" % r] = "=" + norm
        ws["C%d" % r] = "=" + hesaplanan
        ws["C%d" % r].number_format = "0"
        ws["C%d" % r].alignment = Alignment(horizontal="center")
        ws["D%d" % r] = "=" + sonuc
        for sutun in "BCD":
            ws["%s%d" % (sutun, r)].font = KALIN
        ws["F%d" % r] = "=" + bicim
        ws["G%d" % r] = "=" + seri
        ws["H%d" % r] = "=" + beklenen
    sayfa_bicimle(ws)


def hakkinda_sayfasi(wb):
    ws = wb.create_sheet("Hakkında")
    satirlar = [
        ("Konteyner ve AWB numarası kontrolü", BASLIK_FONT),
        ("", None),
        ("Ne yapar", KALIN),
        ("Konteyner numaralarının (ISO 6346) ve hava konşimento numaralarının (AWB) kontrol hanesini formüllerle denetler. "
         "Numaraları sarı hücrelere yapıştırırsınız, sonuç yanındaki sütunlarda görünür.", None),
        ("", None),
        ("Sınırları", KALIN),
        ("Geçerli bir kontrol hanesi yalnızca numaranın biçimce doğru olduğunu gösterir. "
         "Konteynerin var olduğunu, o sahibe ait olduğunu ya da AWB'nin gerçekten düzenlendiğini göstermez.", None),
        ("Her sayfa en çok %d numara denetler. Daha fazlası için formülleri aşağı doğru kopyalayın." % SATIR_SAYISI, None),
        ("Boşluk, tire, nokta ve eğik çizgi silinir. Başka bir karakter varsa numara biçim hatalı sayılır.", None),
        ("Havayolu önekinin gerçek bir havayoluna ait olup olmadığı denetlenmez.", None),
        ("Dosya makro içermez. İndirdiğiniz dosya Korumalı Görünüm'de açılırsa \"Düzenlemeyi Etkinleştir\" düğmesine tıklayın.", None),
        ("Yardımcı hesaplar gizli sütunlardadır. Görmek için sütun başlıklarının üstündeki + işaretine tıklayın.", None),
        ("", None),
        ("Sürüm tarihi: %s" % SURUM_TARIHI, None),
        ("MIT lisansı", None),
    ]
    for i, (metin, font) in enumerate(satirlar, start=1):
        hucre = ws.cell(row=i, column=1, value=metin)
        hucre.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            hucre.font = font
    baglanti = ws.cell(row=len(satirlar) + 1, column=1, value="LogisticsIQ hazırladı: " + SITE)
    baglanti.hyperlink = SITE
    baglanti.font = BAGLANTI_FONT
    ws.column_dimensions["A"].width = 110
    for i, (metin, _) in enumerate(satirlar, start=1):
        if len(metin) > 110:
            ws.row_dimensions[i].height = 32
    ws.freeze_panes = "A2"


def olustur(yol):
    wb = Workbook()
    konteyner_sayfasi(wb)
    awb_sayfasi(wb)
    hakkinda_sayfasi(wb)
    wb.properties.title = "Konteyner ve AWB numarası kontrolü"
    wb.properties.creator = "LogisticsIQ"
    wb.properties.created = datetime.datetime(2026, 10, 19)
    wb.properties.modified = datetime.datetime(2026, 10, 19)
    wb.active = 0
    wb.save(yol)
    return yol


if __name__ == "__main__":
    hedef = os.path.join(os.path.dirname(os.path.abspath(__file__)), "konteyner-awb-kontrol.xlsx")
    print("Yazıldı:", olustur(hedef))
