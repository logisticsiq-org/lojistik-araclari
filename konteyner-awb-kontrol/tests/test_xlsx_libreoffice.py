"""Excel dosyasındaki formüllerin Python referans hesabıyla aynı sonucu verdiğini LibreOffice ile sınar.

Örnek numaralar dosyanın bir kopyasına yazılır, LibreOffice her formülü yeniden hesaplar ve
Konteyner ile AWB sayfalarındaki B, C, D sütunları Python sonucuyla karşılaştırılır.
LibreOffice kurulu değilse test atlanır.
Çalıştırma: python3 -m unittest discover -s konteyner-awb-kontrol/tests -v
"""
import os
import shutil
import sys
import tempfile
import unittest

import openpyxl

ARAC_KLASORU = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEPO_KLASORU = os.path.dirname(ARAC_KLASORU)
sys.path.insert(0, ARAC_KLASORU)
sys.path.insert(0, os.path.join(DEPO_KLASORU, "ortak"))
import kontrol_hanesi as kh  # noqa: E402
from libreoffice_hesapla import LibreOfficeYok, hesapla, hucre, soffice_yolu  # noqa: E402

XLSX = os.path.join(ARAC_KLASORU, "konteyner-awb-kontrol.xlsx")
ILK_SATIR = 5


def gecerli_konteyner(sahip_kategori, altı_rakam):
    no10 = sahip_kategori + altı_rakam
    return no10 + str(kh.konteyner_kontrol_hanesi(no10))


def mod11_on_bulan():
    for sayi in range(1000000):
        no10 = "MSKU%06d" % sayi
        toplam = sum((kh.HARF_DEGERLERI[k] if k.isalpha() else int(k)) * 2 ** i for i, k in enumerate(no10))
        if toplam % 11 == 10:
            return no10 + "0"


# (girdi, yalnız sonuç sütunu karşılaştırılsın mı)
KONTEYNER_ORNEKLERI = [
    ("CSQU3054383", False),                  # bilinen geçerli örnek
    ("csqu 305438 3", False),                # küçük harf ve boşluk
    ("CSQU-305438-3", False),                # tire
    ("CSQU.305438.3", False),                # nokta
    ("CSQU/3054383", False),                 # eğik çizgi
    ("CSQU3054384", False),                  # yanlış kontrol hanesi
    ("CSQU3054380", False),
    ("CSQX3054383", False),                  # yanlış kategori harfi
    ("CSQU305438", False),                   # kısa
    ("CSQU30543830", False),                 # uzun
    ("CSQU30543A3", False),                  # rakam yerine harf
    ("ABC", False),
    (" ", False),                            # yalnız boşluk
    (None, False),                           # boş
    (gecerli_konteyner("MSKU", "123456"), False),
    (gecerli_konteyner("OOLJ", "000001"), False),   # J kategorisi
    (gecerli_konteyner("ABCZ", "999999"), False),   # Z kategorisi
    (gecerli_konteyner("TRIU", "654321").lower(), False),  # küçük i harfi içeren sahip kodu
    (mod11_on_bulan(), False),               # mod 11 sonucu 10, hane 0
    ("CSQUı054383", True),                   # noktasız küçük ı: yalnız sonuç karşılaştırılır
]

AWB_ORNEKLERI = [
    ("020-12345675", False),
    ("02012345675", False),
    ("020 12345675", False),
    ("020.12345675", False),
    ("020/12345675", False),
    ("020-12345676", False),                 # yanlış hane
    ("176-12345670", False),                 # yanlış hane, farklı önek
    ("000-00000000", False),                 # kalan 0
    ("999-99999990", False),
    ("0201234567", False),                   # kısa
    ("020123456750", False),                 # uzun
    ("02A12345675", False),                  # harf
    ("CSQU3054383", False),
    (" ", False),
    (None, False),
    ("125-35800420", False),
]


def beklenen_satir(girdi, tur):
    """(temizlenmiş numara, hesaplanan hane, sonuç metni) döndürür. Boş girdi için üçü de boştur."""
    if girdi is None:
        return "", "", ""
    islev = kh.konteyner_kontrol if tur == "konteyner" else kh.awb_kontrol
    durum, beklenen = islev(girdi)
    hane = "" if beklenen is None else str(beklenen)
    return kh.normalle(girdi), hane, kh.sonuc_metni(durum, beklenen)


class LibreOfficeTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            soffice_yolu()
        except LibreOfficeYok as hata:
            raise unittest.SkipTest(str(hata))
        cls.klasor = tempfile.mkdtemp(prefix="kontrol-xlsx-")
        kopya = os.path.join(cls.klasor, "kopya.xlsx")
        shutil.copyfile(XLSX, kopya)
        wb = openpyxl.load_workbook(kopya)
        for sayfa, ornekler in (("Konteyner", KONTEYNER_ORNEKLERI), ("AWB", AWB_ORNEKLERI)):
            ws = wb[sayfa]
            for i, (girdi, _) in enumerate(ornekler):
                if girdi is not None:
                    ws["A%d" % (ILK_SATIR + i)] = girdi
        wb.save(kopya)
        try:
            cls.sayfalar = hesapla(kopya)
        except LibreOfficeYok as hata:
            raise unittest.SkipTest(str(hata))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.klasor, ignore_errors=True)

    def karsilastir(self, sayfa, ornekler, tur):
        for i, (girdi, yalniz_sonuc) in enumerate(ornekler):
            satir = ILK_SATIR + i
            norm, hane, sonuc = beklenen_satir(girdi, tur)
            with self.subTest(sayfa=sayfa, girdi=girdi):
                gercek_sonuc = hucre(self.sayfalar, sayfa, "D%d" % satir)
                self.assertEqual(gercek_sonuc, sonuc)
                if yalniz_sonuc:
                    continue
                self.assertEqual(hucre(self.sayfalar, sayfa, "B%d" % satir), norm)
                self.assertEqual(hucre(self.sayfalar, sayfa, "C%d" % satir), hane)

    def test_konteyner_sayfasi(self):
        self.karsilastir("Konteyner", KONTEYNER_ORNEKLERI, "konteyner")

    def test_awb_sayfasi(self):
        self.karsilastir("AWB", AWB_ORNEKLERI, "awb")

    def test_bos_satirlar_bos_kalir(self):
        for sayfa, ornekler in (("Konteyner", KONTEYNER_ORNEKLERI), ("AWB", AWB_ORNEKLERI)):
            for satir in (ILK_SATIR + len(ornekler), ILK_SATIR + 500, 1004):
                for sutun in "BCD":
                    self.assertEqual(hucre(self.sayfalar, sayfa, "%s%d" % (sutun, satir)), "")

    def test_sayfa_adlari(self):
        self.assertEqual(set(self.sayfalar), {"Konteyner", "AWB", "Hakkında"})

    def test_hakkinda_icerigi(self):
        metin = "\n".join(" ".join(s) for s in self.sayfalar["Hakkında"])
        self.assertIn("2026-10-19", metin)
        self.assertIn("MIT lisansı", metin)
        self.assertIn("LogisticsIQ hazırladı: https://logisticsiq.org", metin)


class CalismaKitabiYapisiTesti(unittest.TestCase):
    """LibreOffice gerektirmeyen yapı denetimleri."""

    def setUp(self):
        self.wb = openpyxl.load_workbook(XLSX)

    def test_formuller_yasakli_islev_icermez(self):
        yasakli = ("LET(", "LAMBDA(", "XLOOKUP(", "FILTER(", "TEXTJOIN(", "INDIRECT(", "OFFSET(", "TEXT(",
                   "NOW(", "TODAY(", "RAND(")
        for ws in self.wb.worksheets:
            for satir in ws.iter_rows(min_row=1, max_row=6):
                for h in satir:
                    if isinstance(h.value, str) and h.value.startswith("="):
                        for ad in yasakli:
                            self.assertNotIn(ad, h.value.upper(), "%s!%s" % (ws.title, h.coordinate))

    def test_giris_hucreleri_sari_ve_dondurulmus_baslik(self):
        for ad in ("Konteyner", "AWB"):
            ws = self.wb[ad]
            self.assertEqual(ws.freeze_panes, "A5")
            self.assertEqual(ws["A5"].fill.start_color.rgb[-6:], "FFF2CC")
            self.assertEqual(ws["A1004"].fill.start_color.rgb[-6:], "FFF2CC")
            self.assertIsNone(ws["D5"].fill.fill_type)
            self.assertTrue(ws["D5"].font.bold)
            self.assertFalse(ws.protection.sheet)

    def test_kosullu_bicimlendirme_var(self):
        for ad in ("Konteyner", "AWB"):
            kurallar = [k for r in self.wb[ad].conditional_formatting for k in r.rules]
            self.assertEqual(len(kurallar), 2)

    def test_hakkinda_baglantisi(self):
        ws = self.wb["Hakkında"]
        baglantilar = [h.hyperlink.target for satir in ws.iter_rows() for h in satir if h.hyperlink]
        self.assertEqual(baglantilar, ["https://logisticsiq.org"])


if __name__ == "__main__":
    unittest.main()
