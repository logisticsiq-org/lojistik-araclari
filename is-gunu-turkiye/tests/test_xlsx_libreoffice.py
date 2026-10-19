"""Excel dosyasındaki formüller, LibreOffice ile yeniden hesaplanıp Python referansıyla karşılaştırılır.

Her senaryo için `is-gunu-turkiye.xlsx` dosyasının bir kopyasına girdiler yazılır, kopya
LibreOffice'e hesaplatılır ve sonuç hücreleri `is_gunu.py` çıktısıyla karşılaştırılır.
Tarihler yerel ayara bağlı metin olarak değil, gün numarası olarak okunur.
LibreOffice yoksa testler atlanır. Dosya Microsoft Excel'de açılarak denenmemiştir.
"""
import csv
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date, timedelta

import openpyxl

KLASOR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEPO = os.path.dirname(KLASOR)
sys.path.insert(0, KLASOR)
sys.path.insert(0, os.path.join(DEPO, "ortak"))

import is_gunu as ig  # noqa: E402
import olustur_xlsx  # noqa: E402

try:
    from libreoffice_hesapla import LibreOfficeYok, hesapla, soffice_yolu  # noqa: E402
    soffice_yolu()
    LIBREOFFICE_VAR = True
except Exception:  # LibreOfficeYok dahil
    LIBREOFFICE_VAR = False

XLSX = os.path.join(KLASOR, "is-gunu-turkiye.xlsx")
CUMARTESI_PAZAR = olustur_xlsx.CUMARTESI_PAZAR
YALNIZ_PAZAR = olustur_xlsx.YALNIZ_PAZAR
UYARI_YOK = olustur_xlsx.UYARI_YOK

# Okunacak hücreler. Tarih hücreleri sayı olarak okunur.
TARIH_HUCRELERI = ["B11"]
DIGER_HUCRELER = ["B10", "B12", "B13", "B18", "B19", "B20", "B21", "B22", "B23"]
LISTE_SATIRLARI = list(range(27, 27 + olustur_xlsx.LISTE_SATIRI))


def seri_tarih(metin):
    return date(1899, 12, 30) + timedelta(days=int(float(metin)))


def hesapla_senaryo(hafta_sonu, b8, b9, b16, b17):
    """Girdileri bir kopyaya yazar, LibreOffice ile hesaplatır ve okunan hücreleri döndürür."""
    gecici = tempfile.mkdtemp(prefix="isgunu-test-")
    try:
        yol = os.path.join(gecici, "senaryo.xlsx")
        shutil.copy(XLSX, yol)
        wb = openpyxl.load_workbook(yol)
        ws = wb["Hesapla"]
        ws["B4"] = hafta_sonu
        ws["B8"] = b8
        ws["B9"] = b9
        ws["B16"] = b16
        ws["B17"] = b17
        oku = wb.create_sheet("Oku")
        okunacaklar = TARIH_HUCRELERI + DIGER_HUCRELER
        for i, adres in enumerate(okunacaklar, start=1):
            oku.cell(row=i, column=1, value=f"=Hesapla!{adres}").number_format = "0"
        satir = len(okunacaklar) + 1
        for r in LISTE_SATIRLARI:
            oku.cell(row=satir, column=1, value=f"=Hesapla!A{r}")
            oku.cell(row=satir, column=2, value=f"=Hesapla!B{r}").number_format = "0"
            satir += 1
        oku.cell(row=satir, column=1, value=f"=Hesapla!A{27 + olustur_xlsx.LISTE_SATIRI}")
        wb.save(yol)
        sayfalar = hesapla(yol)
        satirlar = sayfalar["Oku"]

        def al(satir, sutun):
            try:
                return satirlar[satir][sutun]
            except IndexError:
                return ""

        sonuc = {adres: al(i, 0) for i, adres in enumerate(okunacaklar)}
        liste = []
        taban = len(okunacaklar)
        for k in range(len(LISTE_SATIRLARI)):
            ad, tarih = al(taban + k, 0), al(taban + k, 1)
            if ad != "" or tarih != "":
                liste.append((seri_tarih(tarih), ad))
        sonuc["liste"] = liste
        sonuc["fazla"] = al(taban + len(LISTE_SATIRLARI), 0)
        return sonuc
    finally:
        shutil.rmtree(gecici, ignore_errors=True)


def ham_tatil_satirlari():
    """CSV'deki kayıtları, birleştirmeden, tarih sırasıyla verir (Excel tablosundaki gibi)."""
    with open(ig.VARSAYILAN_CSV, encoding="utf-8", newline="") as f:
        satirlar = [(date.fromisoformat(s["tarih"]), s["ad"]) for s in csv.DictReader(f)]
    return sorted(satirlar, key=lambda x: x[0])


# (hafta sonu seçeneği, başlangıç, eklenecek gün, ilk tarih, son tarih)
SENARYOLAR = [
    (CUMARTESI_PAZAR, date(2026, 5, 25), 3, date(2026, 5, 1), date(2026, 5, 31)),
    (CUMARTESI_PAZAR, date(2026, 3, 16), 5, date(2026, 3, 14), date(2026, 3, 28)),
    (YALNIZ_PAZAR, date(2026, 3, 14), 5, date(2026, 3, 14), date(2026, 3, 28)),
    (CUMARTESI_PAZAR, date(2026, 6, 2), -3, date(2026, 12, 24), date(2027, 1, 4)),
    (YALNIZ_PAZAR, date(2027, 5, 14), 2, date(2027, 5, 10), date(2027, 5, 21)),
    (CUMARTESI_PAZAR, date(2027, 5, 14), 1, date(2027, 5, 14), date(2027, 5, 14)),
    (CUMARTESI_PAZAR, date(2026, 10, 26), 4, date(2026, 10, 26), date(2026, 10, 30)),
    (CUMARTESI_PAZAR, date(2028, 12, 28), 1, date(2028, 12, 1), date(2028, 12, 31)),
    (CUMARTESI_PAZAR, date(2026, 5, 27), 0, date(2026, 1, 1), date(2028, 12, 31)),
    (YALNIZ_PAZAR, date(2028, 2, 24), 2, date(2028, 2, 20), date(2028, 3, 3)),
    (YALNIZ_PAZAR, date(2027, 3, 5), -4, date(2027, 3, 8), date(2027, 3, 11)),
]


@unittest.skipUnless(LIBREOFFICE_VAR, "LibreOffice bulunamadı, Excel formülleri LibreOffice ile denenemedi")
class FormulPythonKarsilastirmaTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        olustur_xlsx.olustur(XLSX)

    def test_dogru_senaryolar(self):
        for hafta_sonu, bas, n, ilk, son in SENARYOLAR:
            with self.subTest(hafta_sonu=hafta_sonu, bas=bas, n=n, ilk=ilk, son=son):
                cumartesi_tatil = hafta_sonu == CUMARTESI_PAZAR
                r = hesapla_senaryo(hafta_sonu, bas, n, ilk, son)

                # Birinci blok
                self.assertEqual(r["B10"], UYARI_YOK)
                beklenen = ig.is_gunu_ekle(bas, n, cumartesi_tatil)
                self.assertEqual(seri_tarih(r["B11"]), beklenen)
                self.assertEqual(r["B12"], ig.gun_adi(beklenen))
                yarim = [t for t in ig.aradaki_tatiller(min(bas, beklenen), max(bas, beklenen)) if t.sure == "yarim"]
                self.assertEqual(int(r["B13"]), len(yarim))

                # İkinci blok
                self.assertEqual(r["B18"], UYARI_YOK)
                o = ig.ozet(ilk, son, cumartesi_tatil)
                self.assertEqual(int(r["B19"]), o["is_gunu"])
                self.assertEqual(int(r["B20"]), o["takvim_gunu"])
                self.assertEqual(int(r["B21"]), o["tam_gun_tatil"])
                self.assertEqual(int(r["B22"]), o["tatil_nedeniyle_dusen_is_gunu"])
                self.assertEqual(int(r["B23"]), o["yarim_gun_tatil"])

                # Tatil listesi: Excel tabloda her kayıt ayrı satırdır, ilk 12 kayıt gösterilir.
                beklenen_liste = [(t, ad) for t, ad in ham_tatil_satirlari() if ilk <= t <= son]
                self.assertEqual(r["liste"], beklenen_liste[:olustur_xlsx.LISTE_SATIRI])
                if len(beklenen_liste) > olustur_xlsx.LISTE_SATIRI:
                    self.assertTrue(r["fazla"].startswith("Aralıkta"))
                else:
                    self.assertEqual(r["fazla"], "")

    def test_kapsam_disi_senaryolar_sonuc_gostermez(self):
        # (başlangıç, n, ilk, son): Python kapsam hatası vermeli, Excel uyarı verip sonucu boş bırakmalı.
        durumlar = [
            (date(2025, 12, 31), 1, date(2025, 12, 30), date(2026, 1, 5)),
            (date(2028, 12, 29), 5, date(2028, 12, 20), date(2029, 1, 3)),
            (date(2026, 1, 3), -2, date(2029, 1, 1), date(2029, 1, 5)),
        ]
        for bas, n, ilk, son in durumlar:
            with self.subTest(bas=bas, n=n, ilk=ilk, son=son):
                with self.assertRaises(ig.KapsamDisiHata):
                    ig.is_gunu_ekle(bas, n)
                with self.assertRaises(ig.KapsamDisiHata):
                    ig.is_gunu_say(ilk, son)
                r = hesapla_senaryo(CUMARTESI_PAZAR, bas, n, ilk, son)
                self.assertIn("kapsam", r["B10"])
                self.assertNotEqual(r["B10"], UYARI_YOK)
                self.assertIn("kapsam", r["B18"])
                for adres in ("B11", "B12", "B13", "B19", "B20", "B21", "B22", "B23"):
                    self.assertEqual(r[adres], "", adres)
                self.assertEqual(r["liste"], [])

    def test_ters_sira_ve_gecersiz_girdi(self):
        with self.assertRaises(ig.TarihSirasiHatasi):
            ig.is_gunu_say(date(2026, 5, 31), date(2026, 5, 1))
        r = hesapla_senaryo(CUMARTESI_PAZAR, date(2026, 5, 25), 2.5, date(2026, 5, 31), date(2026, 5, 1))
        self.assertEqual(r["B10"], "İş günü sayısı tam sayı olmalı.")
        self.assertEqual(r["B18"], "Son tarih ilk tarihten önce olamaz. Sonuç gösterilmiyor.")
        self.assertEqual(r["B11"], "")
        self.assertEqual(r["B19"], "")
        with self.assertRaises(ValueError):
            ig.is_gunu_ekle(date(2026, 5, 25), 2.5)


@unittest.skipUnless(LIBREOFFICE_VAR, "LibreOffice bulunamadı")
class YapiTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        olustur_xlsx.olustur(XLSX)

    def test_sayfalar_ve_hakkinda_metni(self):
        wb = openpyxl.load_workbook(XLSX)
        self.assertEqual(wb.sheetnames, ["Hesapla", "Tatiller", "Hakkında"])
        metin = "\n".join(str(h.value) for satir in wb["Hakkında"].iter_rows() for h in satir if h.value)
        self.assertIn("2026-10-19", metin)
        self.assertIn("MIT lisansı", metin)
        self.assertIn("LogisticsIQ hazırladı: https://logisticsiq.org", metin)

    def test_kapsam_python_ile_ayni(self):
        ws = openpyxl.load_workbook(XLSX)["Tatiller"]
        self.assertEqual(ws["L2"].value.date(), ig.KAPSAM_BASLANGIC)
        self.assertEqual(ws["L3"].value.date(), ig.KAPSAM_BITIS)

    def test_ayni_girdiden_ayni_dosya_cikar(self):
        gecici = tempfile.mkdtemp(prefix="isgunu-uretim-")
        try:
            a = olustur_xlsx.olustur(os.path.join(gecici, "a.xlsx"))
            b = olustur_xlsx.olustur(os.path.join(gecici, "b.xlsx"))
            with open(a, "rb") as fa, open(b, "rb") as fb:
                self.assertEqual(fa.read(), fb.read())
        finally:
            shutil.rmtree(gecici, ignore_errors=True)

    def test_tatil_tablosu_csv_ile_ayni(self):
        wb = openpyxl.load_workbook(XLSX)
        ws = wb["Tatiller"]
        tablo = []
        for satir in ws.iter_rows(min_row=2, max_col=4, values_only=True):
            if satir[0] is None:
                continue
            tablo.append((satir[0].date(), satir[2], satir[3]))
        with open(ig.VARSAYILAN_CSV, encoding="utf-8", newline="") as f:
            csv_satirlari = [(date.fromisoformat(s["tarih"]), s["ad"], "Tam gün" if s["sure"] == "tam" else "Yarım gün")
                             for s in csv.DictReader(f)]
        self.assertEqual(tablo, sorted(csv_satirlari, key=lambda x: x[0]))


if __name__ == "__main__":
    unittest.main()
