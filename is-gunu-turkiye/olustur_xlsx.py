"""`is-gunu-turkiye.xlsx` dosyasını üretir.

Kullanım:
    python3 olustur_xlsx.py

Dosya yalnız formül içerir, makro yoktur. Tatil verisi `resmi-tatiller-2026-2028.csv`
dosyasından okunur. Aynı girdiden her çalıştırmada aynı dosya çıkar.
"""
import csv
import os
import re
import zipfile
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

KLASOR = os.path.dirname(os.path.abspath(__file__))
CSV_YOLU = os.path.join(KLASOR, "resmi-tatiller-2026-2028.csv")
CIKTI_YOLU = os.path.join(KLASOR, "is-gunu-turkiye.xlsx")

SURUM_TARIHI = "2026-10-19"
SITE = "https://logisticsiq.org"
SON_SATIR = 300          # formüllerin okuduğu son tatil satırı
LISTE_SATIRI = 12        # Hesapla sayfasında gösterilen en fazla tatil sayısı
TAM = "Tam gün"
YARIM = "Yarım gün"
CUMARTESI_PAZAR = "Cumartesi ve Pazar tatil"
YALNIZ_PAZAR = "Yalnız Pazar tatil"
UYARI_YOK = "Uyarı yok."

SARI = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
BASLIK_DOLGU = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
KALIN = Font(bold=True)
GRI = Font(color="808080", italic=True)
TARIH_BICIMI = "dd.mm.yyyy"

GUN_FORMULU = ('CHOOSE(WEEKDAY({h},2),"Pazartesi","Salı","Çarşamba","Perşembe","Cuma","Cumartesi","Pazar")')


def tatilleri_oku():
    satirlar = []
    with open(CSV_YOLU, encoding="utf-8", newline="") as f:
        for s in csv.DictReader(f):
            satirlar.append((
                datetime.strptime(s["tarih"], "%Y-%m-%d").date(),
                s["ad"],
                TAM if s["sure"] == "tam" else YARIM,
                s["kaynak"],
            ))
    return sorted(satirlar, key=lambda x: x[0])


def kalin_yap(hucre):
    hucre.font = KALIN


def baslik_yaz(ws, adres, metin):
    ws[adres] = metin
    ws[adres].font = Font(bold=True, size=12)
    ws[adres].fill = BASLIK_DOLGU


def girdi(ws, adres, deger, bicim=None):
    h = ws[adres]
    h.value = deger
    h.fill = SARI
    if bicim:
        h.number_format = bicim
    return h


def uyari_bicimi(ws, adres):
    kural = FormulaRule(
        formula=[f'{adres}<>"{UYARI_YOK}"'],
        font=Font(bold=True, color="9C0006"),
        fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"),
    )
    ws.conditional_formatting.add(adres.replace("$", ""), kural)


def tatiller_sayfasi(wb, tatiller):
    ws = wb.create_sheet("Tatiller")
    basliklar = ["Tarih", "Gün", "Tatil adı", "Süre", "Kaynak", "", "Tam gün tarihleri (formüller kullanır)",
                 "Yarım gün tarihleri (formüller kullanır)", "Aralıktaki sıra (Hesapla sayfası kullanır)"]
    for i, b in enumerate(basliklar, start=1):
        h = ws.cell(row=1, column=i, value=b)
        if b:
            h.font = KALIN
            h.fill = BASLIK_DOLGU
            h.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45

    bos_yedek = 30
    for sira in range(2, SON_SATIR + 1):
        i = sira - 2
        if i < len(tatiller):
            tarih, ad, sure, kaynak = tatiller[i]
            ws.cell(row=sira, column=1, value=tarih)
            ws.cell(row=sira, column=3, value=ad)
            ws.cell(row=sira, column=4, value=sure)
            ws.cell(row=sira, column=5, value=kaynak)
        ws.cell(row=sira, column=1).number_format = TARIH_BICIMI
        if i < len(tatiller) + bos_yedek:
            for k in (1, 3, 4, 5):
                ws.cell(row=sira, column=k).fill = SARI
        ws.cell(row=sira, column=2, value=f'=IF(A{sira}="","",{GUN_FORMULU.format(h=f"A{sira}")})')
        # İlk geçen kayıt tarihi verir. Aynı güne düşen ikinci kayıt sıfır verir, böylece gün bir kez sayılır.
        ws.cell(row=sira, column=7, value=(
            f'=IF(AND(D{sira}="{TAM}",COUNTIFS($A$2:A{sira},A{sira},$D$2:D{sira},"{TAM}")=1),A{sira},0)'))
        ws.cell(row=sira, column=8, value=(
            f'=IF(AND(D{sira}="{YARIM}",COUNTIFS($A$2:A{sira},A{sira},$D$2:D{sira},"{YARIM}")=1),A{sira},0)'))
        # Sıfır, boş hücre gibi davranır ve ekranda gizlenir.
        ws.cell(row=sira, column=7).number_format = TARIH_BICIMI + ";;"
        ws.cell(row=sira, column=8).number_format = TARIH_BICIMI + ";;"
        ws.cell(row=sira, column=7).font = GRI
        ws.cell(row=sira, column=8).font = GRI
        ws.cell(row=sira, column=9, value=(
            f'=IF(AND(ISNUMBER(A{sira}),Hesapla!$B$18="{UYARI_YOK}",A{sira}>=Hesapla!$B$16,A{sira}<=Hesapla!$B$17),'
            f'MAX(I$1:I{sira - 1})+1,"")'))
        ws.cell(row=sira, column=9).font = GRI

    dv = DataValidation(type="list", formula1=f'"{TAM},{YARIM}"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Süre", error=f"{TAM} veya {YARIM} seçin.")
    ws.add_data_validation(dv)
    dv.add(f"D2:D{SON_SATIR}")

    # Kapsam ve yönergeler
    ws["K1"] = "Kapsam ve yeni yıl ekleme"
    ws["K1"].font = KALIN
    ws["K1"].fill = BASLIK_DOLGU
    ws["L1"].fill = BASLIK_DOLGU
    ws["K2"] = "Kapsamın ilk günü"
    girdi(ws, "L2", date(min(t[0].year for t in tatiller), 1, 1), TARIH_BICIMI)
    ws["K3"] = "Kapsamın son günü"
    girdi(ws, "L3", date(max(t[0].year for t in tatiller), 12, 31), TARIH_BICIMI)
    ws["K4"] = "Kapsam metni (otomatik)"
    ws["L4"] = ('=RIGHT("0"&DAY(L2),2)&"."&RIGHT("0"&MONTH(L2),2)&"."&YEAR(L2)&" ile "'
                '&RIGHT("0"&DAY(L3),2)&"."&RIGHT("0"&MONTH(L3),2)&"."&YEAR(L3)')
    ws["L4"].font = GRI
    yonerge = [
        "Yeni yılın tatilleri Diyanet tarafından yayımlandığında şunları yapın:",
        "1. A, C, D ve E sütunlarında son dolu satırın altına her tatil için bir satır ekleyin. "
        "Tarihi A sütununa, adı C sütununa yazın, D sütunundan süreyi seçin.",
        "2. Bayram arifeleri ve 28 Ekim öğleden sonrası gibi yarım gün tatillerde süre olarak Yarım gün seçin.",
        "3. Aynı güne iki tatil denk gelirse iki satır da yazılabilir. Gün yalnız bir kez sayılır.",
        "4. L3 hücresindeki kapsamın son gününü yeni yılın 31 Aralık tarihine ilerletin.",
        "G, H ve I sütunlarındaki formüllere dokunmayın. Satırlar 300. satıra kadar okunur.",
        "Kapsamın dışındaki tarihlerde Hesapla sayfası sonuç göstermez ve uyarı verir.",
    ]
    for i, m in enumerate(yonerge):
        h = ws.cell(row=6 + i, column=11, value=m)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=6 + i, start_column=11, end_row=6 + i, end_column=12)
        ws.row_dimensions[6 + i].height = 32 if len(m) > 70 else 18

    for col, genislik in zip("ABCDEFGHIJKL", (13, 14, 42, 12, 56, 3, 22, 22, 22, 3, 30, 22)):
        ws.column_dimensions[col].width = genislik
    ws.freeze_panes = "A2"

    wb.defined_names["TamGunTatiller"] = DefinedName("TamGunTatiller", attr_text=f"Tatiller!$G$2:$G${SON_SATIR}")
    wb.defined_names["YarimGunTatiller"] = DefinedName("YarimGunTatiller", attr_text=f"Tatiller!$H$2:$H${SON_SATIR}")
    wb.defined_names["KapsamBas"] = DefinedName("KapsamBas", attr_text="Tatiller!$L$2")
    wb.defined_names["KapsamBit"] = DefinedName("KapsamBit", attr_text="Tatiller!$L$3")
    wb.defined_names["KapsamMetni"] = DefinedName("KapsamMetni", attr_text="Tatiller!$L$4")
    return ws


def hesapla_sayfasi(wb):
    ws = wb.active
    ws.title = "Hesapla"
    ws["A1"] = "İş günü hesaplayıcı (Türkiye)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Sarı hücrelere yazın. Kalın yazılı hücreler sonuçtur. Ayrıntılar için Hakkında sayfasına bakın."
    ws["A2"].font = GRI

    ws["A4"] = "Hafta sonu"
    girdi(ws, "B4", CUMARTESI_PAZAR)
    ws["C4"] = "Cumartesi günü çalışılıyorsa Yalnız Pazar tatil seçeneğini seçin."
    ws["A5"] = "Tatil listesinin kapsamı"
    ws["B5"] = "=KapsamMetni"
    ws["C5"] = "Yarım gün tatiller (bayram arifeleri, 28 Ekim) iş günü sayılır ve ayrıca gösterilir."
    for h in ("C4", "C5"):
        ws[h].font = GRI
    dv = DataValidation(type="list", formula1=f'"{CUMARTESI_PAZAR},{YALNIZ_PAZAR}"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Hafta sonu", error="Listeden bir seçenek seçin.")
    ws.add_data_validation(dv)
    dv.add("B4")

    ws["E3"] = "Yardımcı hücreler (değiştirmeyin)"
    ws["E3"].font = GRI
    ws["E4"] = "Hafta sonu kodu"
    ws["F4"] = f'=IF(B4="{YALNIZ_PAZAR}",11,1)'

    # Blok 1
    baslik_yaz(ws, "A7", "1. Başlangıç tarihine iş günü ekleme")
    ws["B7"].fill = BASLIK_DOLGU
    ws["C7"].fill = BASLIK_DOLGU
    ws["A8"] = "Başlangıç tarihi"
    girdi(ws, "B8", date(2026, 5, 25), TARIH_BICIMI)
    ws["A9"] = "Eklenecek iş günü sayısı"
    girdi(ws, "B9", 3, "0")
    ws["C9"] = "Geriye gitmek için eksi işaretli sayı yazın, örneğin -3. Başlangıç günü sayılmaz."
    ws["C9"].font = GRI
    ws["A10"] = "Uyarı"
    ws["E8"] = "Ham sonuç"
    ws["F8"] = ('=IF(AND(ISNUMBER(B8),ISNUMBER(B9)),IFERROR(WORKDAY.INTL(B8,B9,F4,TamGunTatiller),""),"")')
    ws["F8"].number_format = TARIH_BICIMI
    ws["B10"] = (
        '=IF(OR(B8="",B9=""),"Başlangıç tarihini ve eklenecek iş günü sayısını girin.",'
        'IF(OR(NOT(ISNUMBER(B8)),NOT(ISNUMBER(B9))),"Başlangıç geçerli bir tarih, iş günü sayısı bir sayı olmalı.",'
        'IF(B9<>INT(B9),"İş günü sayısı tam sayı olmalı.",'
        'IF(OR(B8<KapsamBas,B8>KapsamBit),"Başlangıç tarihi kapsam dışında ("&KapsamMetni&"). Sonuç gösterilmiyor.",'
        'IF(OR(F8="",F8<KapsamBas,F8>KapsamBit),"Sonuç kapsam dışına çıkıyor ("&KapsamMetni&"). Sonuç gösterilmiyor.",'
        f'"{UYARI_YOK}")))))'
    )
    ws["A11"] = "Sonuç tarihi"
    ws["B11"] = f'=IF($B$10="{UYARI_YOK}",F8,"")'
    ws["B11"].number_format = TARIH_BICIMI
    ws["A12"] = "Sonuç günü"
    ws["B12"] = f'=IF($B$10="{UYARI_YOK}",{GUN_FORMULU.format(h="F8")},"")'
    ws["A13"] = "Aradaki yarım gün tatil sayısı (iş günü sayıldı)"
    ws["B13"] = (f'=IF($B$10="{UYARI_YOK}",COUNTIFS(YarimGunTatiller,">="&MIN(B8,F8),'
                 f'YarimGunTatiller,"<="&MAX(B8,F8)),"")')
    ws["B13"].number_format = "0"
    for h in ("B11", "B12", "B13"):
        kalin_yap(ws[h])
        ws[h].alignment = Alignment(horizontal="left")
    uyari_bicimi(ws, "B10")

    # Blok 2
    baslik_yaz(ws, "A15", "2. İki tarih arasındaki iş günü")
    ws["B15"].fill = BASLIK_DOLGU
    ws["C15"].fill = BASLIK_DOLGU
    ws["A16"] = "İlk tarih"
    girdi(ws, "B16", date(2026, 5, 1), TARIH_BICIMI)
    ws["A17"] = "Son tarih"
    girdi(ws, "B17", date(2026, 5, 31), TARIH_BICIMI)
    ws["C17"] = "İki tarih de sayıma dahildir."
    ws["C17"].font = GRI
    ws["A18"] = "Uyarı"
    ws["B18"] = (
        '=IF(OR(B16="",B17=""),"İlk ve son tarihi girin.",'
        'IF(OR(NOT(ISNUMBER(B16)),NOT(ISNUMBER(B17))),"İlk ve son tarih geçerli birer tarih olmalı.",'
        'IF(OR(B16<KapsamBas,B16>KapsamBit,B17<KapsamBas,B17>KapsamBit),'
        '"Tarih kapsam dışında ("&KapsamMetni&"). Sonuç gösterilmiyor.",'
        'IF(B17<B16,"Son tarih ilk tarihten önce olamaz. Sonuç gösterilmiyor.",'
        f'"{UYARI_YOK}"))))'
    )
    uyari_bicimi(ws, "B18")
    ok = f'$B$18="{UYARI_YOK}"'
    ws["A19"] = "İş günü sayısı"
    ws["B19"] = f'=IF({ok},NETWORKDAYS.INTL(B16,B17,F4,TamGunTatiller),"")'
    ws["A20"] = "Takvim günü sayısı"
    ws["B20"] = f'=IF({ok},B17-B16+1,"")'
    ws["A21"] = "Tam gün resmî tatil sayısı"
    ws["B21"] = f'=IF({ok},COUNTIFS(TamGunTatiller,">="&B16,TamGunTatiller,"<="&B17),"")'
    ws["A22"] = "Bunlardan iş günü sayısını düşüren"
    ws["C22"] = "Hafta sonuna denk gelen tatiller zaten iş günü olmadığı için düşürmez."
    ws["C22"].font = GRI
    ws["B22"] = (f'=IF({ok},NETWORKDAYS.INTL(B16,B17,F4)-NETWORKDAYS.INTL(B16,B17,F4,TamGunTatiller),"")')
    ws["A23"] = "Yarım gün tatil sayısı (iş günü sayıldı)"
    ws["B23"] = f'=IF({ok},COUNTIFS(YarimGunTatiller,">="&B16,YarimGunTatiller,"<="&B17),"")'
    for h in ("B19", "B20", "B21", "B22", "B23"):
        kalin_yap(ws[h])
        ws[h].number_format = "0"
        ws[h].alignment = Alignment(horizontal="left")

    # Aralıktaki tatiller
    baslik_yaz(ws, "A25", "Aralıktaki resmî tatiller (ikinci bloktaki tarihler)")
    ws["B25"].fill = BASLIK_DOLGU
    ws["C25"].fill = BASLIK_DOLGU
    for adres, metin in (("A26", "Tatil"), ("B26", "Tarih"), ("C26", "Gün ve süre")):
        ws[adres] = metin
        ws[adres].font = KALIN
    ws["F26"] = "Sıra"
    ws["G26"] = "Tatiller sayfasındaki satır"
    ws["F26"].font = GRI
    ws["G26"].font = GRI
    for k in range(1, LISTE_SATIRI + 1):
        r = 26 + k
        ws[f"F{r}"] = k
        ws[f"G{r}"] = f'=IFERROR(MATCH(F{r},Tatiller!$I$2:$I${SON_SATIR},0),"")'
        ws[f"A{r}"] = f'=IF(G{r}="","",INDEX(Tatiller!$C$2:$C${SON_SATIR},G{r}))'
        ws[f"B{r}"] = f'=IF(G{r}="","",INDEX(Tatiller!$A$2:$A${SON_SATIR},G{r}))'
        ws[f"B{r}"].number_format = TARIH_BICIMI
        ws[f"B{r}"].alignment = Alignment(horizontal="left")
        ws[f"C{r}"] = (f'=IF(G{r}="","",INDEX(Tatiller!$B$2:$B${SON_SATIR},G{r})&", "&'
                       f'IF(INDEX(Tatiller!$D$2:$D${SON_SATIR},G{r})="{YARIM}","yarım gün, iş günü sayılır","tam gün"))')
        ws[f"F{r}"].font = GRI
        ws[f"G{r}"].font = GRI
    son = 26 + LISTE_SATIRI + 1
    ws[f"A{son}"] = (f'=IF(MAX(Tatiller!$I$2:$I${SON_SATIR})>{LISTE_SATIRI},'
                     f'"Aralıkta "&MAX(Tatiller!$I$2:$I${SON_SATIR})&" tatil kaydı var. '
                     f'Tamamı için Tatiller sayfasına bakın.","")')
    ws[f"A{son}"].font = GRI

    for col, genislik in zip("ABCDEFG", (46, 30, 78, 3, 24, 12, 26)):
        ws.column_dimensions[col].width = genislik
    ws.freeze_panes = "A3"
    return ws


def hakkinda_sayfasi(wb):
    ws = wb.create_sheet("Hakkında")
    satirlar = [
        ("İş günü hesaplayıcı (Türkiye)", "baslik"),
        ("", None),
        ("Ne yapar", "alt"),
        ("Bu dosya bir tarihe iş günü ekler veya çıkarır, iki tarih arasındaki iş günlerini sayar. "
         "Hesapta Türkiye'deki resmî tatiller dikkate alınır.", None),
        ("", None),
        ("Nasıl kullanılır", "alt"),
        ("1. İndirdiğiniz dosya Korumalı Görünüm'de açılırsa Düzenlemeyi Etkinleştir düğmesine tıklayın.", None),
        ("2. Hesapla sayfasında hafta sonu seçeneğini kontrol edin. Cumartesi çalışılıyorsa Yalnız Pazar tatil seçin.", None),
        ("3. Birinci blokta başlangıç tarihini ve eklenecek iş günü sayısını yazın.", None),
        ("4. İkinci blokta iki tarih yazın. İki tarih de sayıma dahildir.", None),
        ("5. Uyarı satırında bir mesaj varsa sonuç gösterilmez. Mesajı okuyup girdiyi düzeltin.", None),
        ("", None),
        ("Sınırlar", "alt"),
        ("Tatil listesi 1 Ocak 2026 ile 31 Aralık 2028 arasını kapsar. Bu aralığın dışında sonuç gösterilmez.", None),
        ("Tam gün tatiller iş günü sayılmaz. Bayram arifeleri ve 28 Ekim öğleden sonrası yarım gün tatildir. "
         "Sabah çalışıldığı için bu günler iş günü sayılır ve ayrıca gösterilir.", None),
        ("Listede yalnızca 2429 sayılı kanuna dayanan resmî tatiller vardır. İdari izinler, köprü günleri ve "
         "işyerine özel tatiller yoktur.", None),
        ("Tatil tarihleri her yıl Diyanet tarafından yayımlanır ve değişebilir. Liste her yıl güncellenmelidir. "
         "Nasıl ekleneceği Tatiller sayfasının sağ tarafında yazıyor.", None),
        ("Sayım kuralı WORKDAY.INTL ve NETWORKDAYS.INTL ile aynıdır. Sayı 0 ise başlangıç tarihi aynen döner.", None),
        ("Bu dosya makro içermez. Sayfalar korumasızdır, bu yüzden formül hücrelerini yanlışlıkla silmemeye dikkat edin.", None),
        ("", None),
        ("Kaynak", "alt"),
        ("Diyanet İşleri Başkanlığı resmî tatil listeleri, 2429 sayılı kanuna göre hazırlanmıştır:", None),
        ("2026: https://vakithesaplama.diyanet.gov.tr/icerik.php?icerik=158", None),
        ("2027: https://vakithesaplama.diyanet.gov.tr/icerik.php?icerik=159", None),
        ("2028: https://vakithesaplama.diyanet.gov.tr/icerik.php?icerik=193", None),
        ("", None),
        (f"Sürüm tarihi: {SURUM_TARIHI}", None),
        ("MIT lisansı", None),
        ("LogisticsIQ hazırladı: https://logisticsiq.org", "bag"),
    ]
    for i, (metin, tur) in enumerate(satirlar, start=1):
        h = ws.cell(row=i, column=1, value=metin if metin else None)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        if tur == "baslik":
            h.font = Font(bold=True, size=14)
        elif tur == "alt":
            h.font = Font(bold=True, size=12)
            h.fill = BASLIK_DOLGU
        elif tur == "bag":
            h.hyperlink = SITE
            h.style = "Hyperlink"
        if metin and len(metin) > 110:
            ws.row_dimensions[i].height = 32
    ws.column_dimensions["A"].width = 115
    return ws


def zaman_damgalarini_sabitle(yol):
    """openpyxl kayıt anındaki saati dosyaya yazar. Aynı girdiden aynı dosya çıksın diye sabitlenir."""
    gecici = yol + ".tmp"
    with zipfile.ZipFile(yol) as kaynak, zipfile.ZipFile(gecici, "w", zipfile.ZIP_DEFLATED) as hedef:
        for bilgi in kaynak.infolist():
            veri = kaynak.read(bilgi.filename)
            if bilgi.filename == "docProps/core.xml":
                veri = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-19T12:00:00Z", veri)
            yeni = zipfile.ZipInfo(bilgi.filename, date_time=(2026, 10, 19, 12, 0, 0))
            yeni.compress_type = zipfile.ZIP_DEFLATED
            yeni.external_attr = bilgi.external_attr
            hedef.writestr(yeni, veri)
    os.replace(gecici, yol)


def olustur(yol=CIKTI_YOLU):
    tatiller = tatilleri_oku()
    wb = Workbook()
    hesapla_sayfasi(wb)
    tatiller_sayfasi(wb, tatiller)
    hakkinda_sayfasi(wb)
    wb.properties.title = "İş günü hesaplayıcı (Türkiye)"
    wb.properties.creator = "LogisticsIQ"
    sabit = datetime(2026, 10, 19, 12, 0, 0)
    wb.properties.created = sabit
    wb.properties.modified = sabit
    wb.save(yol)
    zaman_damgalarini_sabitle(yol)
    return yol


if __name__ == "__main__":
    print(olustur())
