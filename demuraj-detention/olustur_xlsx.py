"""`demuraj-detention.xlsx` dosyasını üretir.

Kullanım:
    python3 olustur_xlsx.py

Dosya yalnız formül içerir, makro yoktur. Örnek girdiler `ornek.json` dosyasından alınır.
Aynı girdiden her çalıştırmada aynı dosya çıkar.
"""
import json
import os
import re
import sys
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

KLASOR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, KLASOR)
import demuraj as d  # noqa: E402

ORNEK_YOLU = os.path.join(KLASOR, "ornek.json")
CIKTI_YOLU = os.path.join(KLASOR, "demuraj-detention.xlsx")
SITE = "https://logisticsiq.org"

SARI = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
BASLIK_DOLGU = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
KALIN = Font(bold=True)
GRI = Font(color="808080", italic=True)
TARIH_BICIMI = "dd.mm.yyyy"
PARA_BICIMI = "#,##0.00"

MOD_METNI = {v: k for k, v in d.MOD_METINLERI.items()}
MOD_AYRI_METNI = MOD_METNI[d.MOD_AYRI]
MOD_BIRLESIK_METNI = MOD_METNI[d.MOD_BIRLESIK]

# Test dosyası da kullanır: tablo adı -> başlık satırı. Satırlar başlık + 2 ile başlık + 6 arasındadır.
TABLO_BASLIKLARI = {"demuraj": 34, "detention": 43, "birlesik": 52, "ardiye": 61}
TABLO_ADLARI = {
    "demuraj": "Demuraj tarifesi (tahliyeden çıkışa kadar)",
    "detention": "Detention tarifesi (çıkıştan boş iadeye kadar)",
    "birlesik": "Birleşik tarife (tahliyeden boş iadeye kadar)",
    "ardiye": "Ardiye tarifesi (tahliyeden çıkışa kadar, isteğe bağlı)",
}
HUCRELER = {
    "konteyner_sayisi": "B5", "konteyner_tipi": "B6", "para_birimi": "B7", "mod": "B8",
    "ilk_gun": "B9", "kur": "B10", "tahliye": "B13", "cikis": "B14", "bos_iade": "B15",
    "uyari": "B16", "demuraj_gun": "B19", "detention_gun": "B20", "birlesik_gun": "B21", "ardiye_gun": "B22",
    "toplam": "B27", "toplam_try": "B28", "cumle": "B29",
}


def satirlar(tablo):
    r0 = TABLO_BASLIKLARI[tablo]
    return list(range(r0 + 2, r0 + 2 + d.SATIR_SAYISI))


def ara_toplam_hucresi(tablo):
    return f"F{TABLO_BASLIKLARI[tablo] + 7}"


def girdi(ws, adres, deger, bicim=None):
    h = ws[adres]
    h.value = deger
    h.fill = SARI
    if bicim:
        h.number_format = bicim
    return h


def baslik(ws, satir, metin, son_sutun="G"):
    ws[f"A{satir}"] = metin
    for sutun in "ABCDEFG"[: "ABCDEFG".index(son_sutun) + 1]:
        ws[f"{sutun}{satir}"].fill = BASLIK_DOLGU
    ws[f"A{satir}"].font = Font(bold=True, size=12)


def zincir(kosullar, varsayilan):
    """(koşul, mesaj) çiftlerinden iç içe IF formülü kurar."""
    sonuc = f'"{varsayilan}"'
    for kosul, mesaj in reversed(kosullar):
        sonuc = f'IF({kosul},"{mesaj}",{sonuc})'
    return sonuc


def uyari_formulu():
    ard = satirlar("ardiye")
    return "=" + zincir([
        ('IF(ISNUMBER(B5),OR(B5<1,B5<>INT(B5)),TRUE)', d.M_SAYI),
        ('AND(B7<>"USD",B7<>"EUR",B7<>"TRY")', d.M_PARA),
        ('$J$5=0', d.M_MOD),
        ('AND(B9<>"Evet",B9<>"Hayır")', d.M_ILK_GUN),
        ('NOT(ISNUMBER(B13))', d.M_TAHLIYE),
        ('AND($J$5=1,NOT(ISNUMBER(B14)))', d.M_CIKIS),
        ('NOT(ISNUMBER(B15))', d.M_IADE),
        ('AND($J$5=2,$J$7=1,NOT(ISNUMBER(B14)))', d.M_ARDIYE_CIKIS),
        ('AND(ISNUMBER(B14),B14<B13)', d.M_SIRA_CIKIS),
        ('B15<B13', d.M_SIRA_IADE_TAHLIYE),
        ('AND(ISNUMBER(B14),B15<B14)', d.M_SIRA_IADE_CIKIS),
        ('AND(B7<>"TRY",B10<>"",IF(ISNUMBER(B10),B10<=0,TRUE))', d.M_KUR),
        ('$J$8>0', d.M_TARIFE),
    ], d.UYARI_YOK)


def satir_notu_formulu(r, uygulanma):
    """Bir tarife satırı için not formülü. `uygulanma` bu tablo bu modda kullanılmıyorsa True veren koşuldur."""
    ic = (f'IF(NOT(ISNUMBER(B{r})),"{d.R_BASLANGIC}",IF(OR(B{r}<1,B{r}<>INT(B{r})),"{d.R_BASLANGIC}",'
          f'IF(NOT(ISNUMBER(D{r})),"{d.R_UCRET_YOK}",IF(D{r}<0,"{d.R_UCRET_EKSI}",'
          f'IF(C{r}="","",IF(NOT(ISNUMBER(C{r})),"{d.R_BITIS_SAYI}",IF(C{r}<>INT(C{r}),"{d.R_BITIS_SAYI}",'
          f'IF(C{r}<B{r},"{d.R_BITIS_KUCUK}",""))))))))')
    ic = f'IF(B{r}="","",{ic})'
    return f'=IF({uygulanma},"",{ic})' if uygulanma else f"={ic}"


def tablo_yaz(ws, tablo, g):
    r0 = TABLO_BASLIKLARI[tablo]
    ok = f'$B$16="{d.UYARI_YOK}"'
    baslik(ws, r0, TABLO_ADLARI[tablo])
    ws[f"E{r0}"] = "Süre (gün)"
    ws[f"E{r0}"].font = KALIN
    gun_formulu = {
        "demuraj": f"=IF(AND({ok},$J$5=1),$B$14-$B$13+$J$6,\"\")",
        "detention": f"=IF(AND({ok},$J$5=1),$B$15-$B$14+$J$6,\"\")",
        "birlesik": f"=IF(AND({ok},$J$5=2),$B$15-$B$13+$J$6,\"\")",
        "ardiye": f"=IF(AND({ok},$J$7=1),$B$14-$B$13+$J$6,\"\")",
    }[tablo]
    ws[f"F{r0}"] = gun_formulu
    ws[f"F{r0}"].font = KALIN
    ws[f"F{r0}"].number_format = "0"
    uygulanmaz = {"demuraj": "$J$5<>1", "detention": "$J$5<>1", "birlesik": "$J$5<>2", "ardiye": None}[tablo]
    if tablo in ("demuraj", "detention"):
        ws[f"C{r0}"] = '=IF($J$5=2,"Bu hesap şeklinde kullanılmaz.","")'
    elif tablo == "birlesik":
        ws[f"C{r0}"] = '=IF($J$5=1,"Bu hesap şeklinde kullanılmaz.","")'
    else:
        ws[f"C{r0}"] = "Boş bırakırsanız ardiye hesaplanmaz."
    ws[f"C{r0}"].font = GRI
    for k, baslik_metni in enumerate(["Satır", "Başlangıç günü", "Bitiş günü (boşsa sonrası)",
                                      "Günlük ücret (konteyner başına)", "Gün sayısı", "Tutar", "Satır notu"]):
        h = ws.cell(row=r0 + 1, column=k + 1, value=baslik_metni)
        h.font = KALIN
        h.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r0 + 1].height = 32
    veri = d._satirlari_duzenle(getattr(g, tablo)) if g else []
    for i, r in enumerate(satirlar(tablo)):
        ws[f"A{r}"] = f"Satır {i + 1}"
        bas, bit, ucret = veri[i] if i < len(veri) else (None, None, None)
        girdi(ws, f"B{r}", bas, "0")
        girdi(ws, f"C{r}", bit, "0")
        girdi(ws, f"D{r}", ucret, PARA_BICIMI)
        ws[f"E{r}"] = (f'=IF(OR(B{r}="",$F${r0}=""),"",MAX(0,MIN($F${r0},IF(C{r}="",$F${r0},C{r}))-B{r}+1))')
        ws[f"E{r}"].number_format = "0"
        ws[f"F{r}"] = f'=IF(E{r}="","",E{r}*D{r}*$B$5)'
        ws[f"F{r}"].number_format = PARA_BICIMI
        ws[f"G{r}"] = satir_notu_formulu(r, uygulanmaz)
    ilk, son = satirlar(tablo)[0], satirlar(tablo)[-1]
    ws[f"A{r0 + 7}"] = "Ara toplam"
    ws[f"A{r0 + 7}"].font = KALIN
    ws[f"F{r0 + 7}"] = f'=IF($F${r0}="","",SUM(F{ilk}:F{son}))'
    ws[f"F{r0 + 7}"].font = KALIN
    ws[f"F{r0 + 7}"].number_format = PARA_BICIMI
    ws.conditional_formatting.add(f"G{ilk}:G{son}", FormulaRule(
        formula=[f"LEN(G{ilk})>0"], font=Font(bold=True, color="9C0006")))
    # Veri doğrulama
    dv_gun = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                            showErrorMessage=True, errorTitle="Gün", error="1 veya daha büyük bir tam sayı yazın.")
    dv_ucret = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Ücret", error="0 veya daha büyük bir sayı yazın.")
    for dv in (dv_gun, dv_ucret):
        ws.add_data_validation(dv)
    dv_gun.add(f"B{ilk}:C{son}")
    dv_ucret.add(f"D{ilk}:D{son}")


def hesapla_sayfasi(wb, g):
    ws = wb.active
    ws.title = "Hesapla"
    ws["A1"] = "Demuraj, detention ve ardiye maliyeti (tahmin)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Bu sonuç bir tahmindir. Sözleşmenizdeki ve terminal tarifenizdeki koşullar geçerlidir. "
                "Tüm günler takvim günüdür.")
    ws["A2"].font = Font(bold=True, color="9C0006")
    ws["A3"] = "Dosyadaki tarihler ve tarifeler örnektir. Sarı hücrelere kendi bilgilerinizi yazın."
    ws["A3"].font = GRI

    baslik(ws, 4, "Ayarlar", "C")
    ayarlar = [
        (5, "Konteyner sayısı", g.konteyner_sayisi, "0", "Tutarlar bu sayıyla çarpılır."),
        (6, "Konteyner tipi (serbest metin)", g.konteyner_tipi, None, "Yalnız özet cümlesinde görünür."),
        (7, "Para birimi", g.para_birimi, None, "Tarifedeki günlük ücretler bu para birimindedir."),
        (8, "Hesap şekli", MOD_METNI[g.mod], None,
         "Taşıyıcınız demuraj ve detention için ayrı süre veriyorsa Ayrı, tek süre veriyorsa Birleşik seçin."),
        (9, "İlk gün sayılır mı?", "Evet" if g.ilk_gun_sayilir else "Hayır", None,
         "Evet: başlangıç günü 1. gündür (gün = bitiş eksi başlangıç artı 1). Hayır: gün = bitiş eksi başlangıç."),
        (10, "TRY kuru (isteğe bağlı)", g.kur_try, "0.0000",
         "1 birim para biriminin TRY karşılığı. Güncel kur için bu depodaki ../tcmb-kurlari/ aracına bakabilirsiniz."),
    ]
    for satir, etiket, deger, bicim, not_ in ayarlar:
        ws[f"A{satir}"] = etiket
        girdi(ws, f"B{satir}", deger, bicim)
        ws[f"C{satir}"] = not_
        ws[f"C{satir}"].font = GRI

    baslik(ws, 12, "Tarihler", "C")
    tarihler = [
        (13, "Tahliye tarihi", g.tahliye, "Konteynerin gemiden indirildiği gün."),
        (14, "Çıkış tarihi", g.cikis, "Konteynerin terminalden çıktığı gün."),
        (15, "Boş iade tarihi", g.bos_iade, "Boş konteynerin iade edildiği gün."),
    ]
    for satir, etiket, deger, not_ in tarihler:
        ws[f"A{satir}"] = etiket
        girdi(ws, f"B{satir}", deger, TARIH_BICIMI)
        ws[f"C{satir}"] = not_
        ws[f"C{satir}"].font = GRI
    ws["A16"] = "Uyarı"
    ws["B16"] = uyari_formulu()
    ws.conditional_formatting.add("B16", FormulaRule(
        formula=[f'$B$16<>"{d.UYARI_YOK}"'], font=Font(bold=True, color="9C0006"),
        fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))

    baslik(ws, 18, "Sonuç (tahmin)", "C")
    sonuc_satirlari = [
        (19, "Demuraj süresi (gün)", f"=F{TABLO_BASLIKLARI['demuraj']}", "0"),
        (20, "Detention süresi (gün)", f"=F{TABLO_BASLIKLARI['detention']}", "0"),
        (21, "Birleşik süre (gün)", f"=F{TABLO_BASLIKLARI['birlesik']}", "0"),
        (22, "Ardiye süresi (gün)", f"=F{TABLO_BASLIKLARI['ardiye']}", "0"),
        (23, "Demuraj tutarı", f"={ara_toplam_hucresi('demuraj')}", PARA_BICIMI),
        (24, "Detention tutarı", f"={ara_toplam_hucresi('detention')}", PARA_BICIMI),
        (25, "Birleşik süre tutarı", f"={ara_toplam_hucresi('birlesik')}", PARA_BICIMI),
        (26, "Ardiye tutarı", f"={ara_toplam_hucresi('ardiye')}", PARA_BICIMI),
    ]
    for satir, etiket, formul, bicim in sonuc_satirlari:
        ws[f"A{satir}"] = etiket
        ws[f"B{satir}"] = formul
        ws[f"B{satir}"].number_format = bicim
        ws[f"B{satir}"].alignment = Alignment(horizontal="left")
    ok = f'$B$16="{d.UYARI_YOK}"'
    ws["A27"] = "Genel toplam (seçilen para birimi)"
    ws["B27"] = (f'=IF({ok},SUM({ara_toplam_hucresi("demuraj")},{ara_toplam_hucresi("detention")},'
                 f'{ara_toplam_hucresi("birlesik")},{ara_toplam_hucresi("ardiye")}),"")')
    ws["A28"] = "Genel toplam (TRY)"
    ws["B28"] = f'=IF({ok},IF($B$7="TRY",B27,IF(ISNUMBER($B$10),B27*$B$10,"")),"")'
    ws["C28"] = "Kur girilmediyse boş kalır."
    ws["C28"].font = GRI
    for h in ("B27", "B28"):
        ws[h].font = Font(bold=True, size=12)
        ws[h].number_format = PARA_BICIMI
        ws[h].alignment = Alignment(horizontal="left")
    ws["A29"] = "Özet"
    ws["A29"].alignment = Alignment(vertical="top")
    ws["B29"] = (
        f'=IF({ok},B5&" adet"&IF(TRIM(B6)="",""," "&TRIM(B6))&" konteyner için tahmini toplam "'
        f'&FIXED(B27,2,TRUE)&" "&B7&". "'
        f'&IF($J$5=1,"Demuraj süresi "&B19&" gün, detention süresi "&B20&" gün","Birleşik süre "&B21&" gün")'
        f'&IF(B22<>"",", ardiye süresi "&B22&" gün","")&"."'
        f'&IF(AND(B7<>"TRY",ISNUMBER(B28))," Verilen kurla yaklaşık "&FIXED(B28,2,TRUE)&" TRY.","")'
        f'&" Bu bir tahmindir, sözleşmenizdeki ve terminal tarifenizdeki koşullar geçerlidir.","")'
    )
    ws["B29"].font = KALIN
    ws["B29"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B29:G29")
    ws.row_dimensions[29].height = 64

    baslik(ws, 31, "Tarifeler")
    ws["A32"] = ("Tarifenizi olduğu gibi girin. Serbest süre, ücreti 0 olan bir satırdır. Bitiş gününü boş "
                 "bırakırsanız satır sonuna kadar sürer. Satırlar çakışmamalıdır.")
    ws["A32"].font = GRI
    for tablo in d.TABLOLAR:
        tablo_yaz(ws, tablo, g)

    # Yardımcı hücreler
    ws["I4"] = "Yardımcı hücreler (değiştirmeyin)"
    ws["I4"].font = GRI
    ard = satirlar("ardiye")
    yardimci = [
        (5, "Hesap şekli kodu", f'=IF(B8="{MOD_AYRI_METNI}",1,IF(B8="{MOD_BIRLESIK_METNI}",2,0))'),
        (6, "İlk gün eki", '=IF(B9="Evet",1,0)'),
        (7, "Ardiye kullanılıyor mu", f"=IF(COUNTA(B{ard[0]}:B{ard[-1]})>0,1,0)"),
        (8, "Hatalı tarife satırı sayısı", "=" + "+".join(
            f"SUMPRODUCT(--(LEN(G{satirlar(t)[0]}:G{satirlar(t)[-1]})>0))" for t in d.TABLOLAR)),
    ]
    for satir, etiket, formul in yardimci:
        ws[f"I{satir}"] = etiket
        ws[f"J{satir}"] = formul
        ws[f"I{satir}"].font = GRI
        ws[f"J{satir}"].font = GRI

    # Açılır listeler ve doğrulamalar
    kurallar = [
        ("B5", DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=False,
                              showErrorMessage=True, errorTitle="Konteyner sayısı",
                              error="1 veya daha büyük bir tam sayı yazın.")),
        ("B7", DataValidation(type="list", formula1='"USD,EUR,TRY"', allow_blank=False,
                              showErrorMessage=True, errorTitle="Para birimi", error="Listeden seçin.")),
        ("B8", DataValidation(type="list", formula1=f'"{MOD_AYRI_METNI},{MOD_BIRLESIK_METNI}"', allow_blank=False,
                              showErrorMessage=True, errorTitle="Hesap şekli", error="Listeden seçin.")),
        ("B9", DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=False,
                              showErrorMessage=True, errorTitle="İlk gün", error="Evet veya Hayır seçin.")),
        ("B10", DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                               showErrorMessage=True, errorTitle="Kur", error="0'dan büyük bir sayı yazın.")),
        ("B13:B15", DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                                   showErrorMessage=True, errorTitle="Tarih", error="Geçerli bir tarih yazın.")),
    ]
    for aralik, dv in kurallar:
        ws.add_data_validation(dv)
        dv.add(aralik)

    for sutun, genislik in zip("ABCDEFGHIJ", (44, 22, 26, 22, 12, 18, 54, 3, 30, 10)):
        ws.column_dimensions[sutun].width = genislik
    ws.freeze_panes = "A4"
    return ws


def hakkinda_sayfasi(wb):
    ws = wb.create_sheet("Hakkında")
    satirlar_ = [
        ("Demuraj, detention ve ardiye maliyeti (tahmin)", "baslik"),
        ("", None),
        ("Ne yapar", "alt"),
        ("Bu dosya, konteynerin tahliye, çıkış ve boş iade tarihlerinden demuraj, detention ve ardiye "
         "maliyetini tahmin eder. Siz taşıyıcınızın ve terminalinizin tarifesini kademe kademe girersiniz.", None),
        ("Sonuç bir tahmindir. Sözleşmenizdeki ve terminal tarifenizdeki koşullar geçerlidir.", None),
        ("", None),
        ("Tanımlar", "alt"),
        ("Demuraj: Dolu konteynerin tahliyeden terminal çıkışına kadar terminalde kaldığı günler.", None),
        ("Detention: Konteynerin terminalden çıkışından boş iadeye kadar terminal dışında geçen günler.", None),
        ("Birleşik süre: Bazı taşıyıcılar tahliyeden boş iadeye kadar tek bir serbest süre uygular. "
         "Bu durumda Hesap şekli olarak Birleşik seçilir.", None),
        ("Ardiye: Terminalin kendi tarifesiyle, demurajla aynı dönem için kestiği depolama bedeli. "
         "Tarife tablosu boşsa ardiye hesaplanmaz.", None),
        ("", None),
        ("Nasıl kullanılır", "alt"),
        ("1. İndirdiğiniz dosya Korumalı Görünüm'de açılırsa Düzenlemeyi Etkinleştir düğmesine tıklayın.", None),
        ("2. Hesapla sayfasında ayarları ve üç tarihi girin. Dosyadaki değerler örnektir.", None),
        ("3. Tarife tablolarını sözleşmenizdeki gibi doldurun. Serbest süre, ücreti 0 olan bir satırdır. "
         "Örneğin 1 ile 7. gün için 0, 8 ile 14. gün için 40, 15. gün ve sonrası için 80.", None),
        ("4. Uyarı satırında bir mesaj varsa sonuç gösterilmez. Mesajı okuyup girdiyi düzeltin.", None),
        ("5. Sonuç bölümünde süreleri, tutarları, genel toplamı ve özet cümlesini görürsünüz.", None),
        ("", None),
        ("Hesap kuralı", "alt"),
        ("Bir satırdaki gün sayısı: en büyük değer (0, en küçük değer (toplam gün, bitiş günü) eksi "
         "başlangıç günü artı 1). Tutar: gün sayısı x günlük ücret x konteyner sayısı.", None),
        ("Ayrı modda demuraj tahliyeden çıkışa, detention çıkıştan boş iadeye kadar hesaplanır. "
         "Birleşik modda tahliyeden boş iadeye kadar tek tablo kullanılır. Ardiye her iki modda tahliyeden çıkışa kadardır.", None),
        ("İlk gün sayılır mı sorusu, başlangıç gününün 1. gün olup olmadığını belirler. Taşıyıcınızın tarifesinde "
         "yazan kurala göre seçin.", None),
        ("İlk gün sayılıyorsa çıkış günü hem demurajın son günü hem detention'ın ilk günü olur. Taşıyıcınız bu günü "
         "yalnızca bir tarafta sayıyorsa sonucu buna göre düzeltin.", None),
        ("", None),
        ("Sınırlar", "alt"),
        ("Tüm sayımlar takvim günüdür. Hafta sonu ve tatiller sayıma dahildir. Hafta sonunu veya tatili "
         "saymayan tarifeler için bu dosya uygun değildir.", None),
        ("Tarifeler taşıyıcıya, limana ve sözleşmeye göre değişir. Bu dosya vergi, KDV, işlem ücreti, "
         "serbest süre uzatması, azami tutar ve konteyner tipine göre farkları kendiliğinden bilmez.", None),
        ("Tarife satırlarının çakışmamasını siz sağlarsınız. Çakışan satırlar iki kez ücretlendirilir.", None),
        ("Kur kutusu isteğe bağlıdır ve elle girilir. Dosya kur çekmez.", None),
        ("Bu dosya makro içermez. Sayfalar korumasızdır, bu yüzden formül hücrelerini yanlışlıkla silmemeye dikkat edin.", None),
        ("", None),
        (f"Sürüm tarihi: {d.SURUM_TARIHI}", None),
        ("MIT lisansı", None),
        ("LogisticsIQ hazırladı: https://logisticsiq.org", "bag"),
    ]
    for i, (metin, tur) in enumerate(satirlar_, start=1):
        h = ws.cell(row=i, column=1, value=metin if metin else None)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        if tur == "baslik":
            h.font = Font(bold=True, size=14)
        elif tur == "alt":
            h.font = Font(bold=True, size=12)
            h.fill = BASLIK_DOLGU
        elif tur == "bag":
            h.hyperlink = SITE
            h.style = "Hyperlink"
        if metin and len(metin) > 110:
            ws.row_dimensions[i].height = 32
    ws.column_dimensions["A"].width = 115
    return ws


def zaman_damgalarini_sabitle(yol):
    """openpyxl kayıt anındaki saati dosyaya yazar. Aynı girdiden aynı dosya çıksın diye sabitlenir."""
    gecici = yol + ".tmp"
    with zipfile.ZipFile(yol) as kaynak, zipfile.ZipFile(gecici, "w", zipfile.ZIP_DEFLATED) as hedef:
        for bilgi in kaynak.infolist():
            veri = kaynak.read(bilgi.filename)
            if bilgi.filename == "docProps/core.xml":
                veri = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-19T12:00:00Z", veri)
            yeni = zipfile.ZipInfo(bilgi.filename, date_time=(2026, 10, 19, 12, 0, 0))
            yeni.compress_type = zipfile.ZIP_DEFLATED
            yeni.external_attr = bilgi.external_attr
            hedef.writestr(yeni, veri)
    os.replace(gecici, yol)


def olustur(yol=CIKTI_YOLU):
    with open(ORNEK_YOLU, encoding="utf-8") as f:
        g = d.girdi_oku(json.load(f))
    wb = Workbook()
    hesapla_sayfasi(wb, g)
    hakkinda_sayfasi(wb)
    wb.properties.title = "Demuraj, detention ve ardiye maliyeti (tahmin)"
    wb.properties.creator = "LogisticsIQ"
    sabit = datetime(2026, 10, 19, 12, 0, 0)
    wb.properties.created = sabit
    wb.properties.modified = sabit
    wb.save(yol)
    zaman_damgalarini_sabitle(yol)
    return yol


if __name__ == "__main__":
    print(olustur())
