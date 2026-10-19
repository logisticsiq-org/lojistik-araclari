"""Excel dosyasındaki formüller, LibreOffice ile yeniden hesaplanıp Python referansıyla karşılaştırılır.

Her senaryo için `demuraj-detention.xlsx` dosyasının bir kopyasına girdiler yazılır, kopya LibreOffice'e
hesaplatılır ve sonuç hücreleri `demuraj.py` çıktısıyla karşılaştırılır. Hatalı girdilerde Uyarı hücresinin
metni Python'un hata metniyle aynı olmalı, sonuç hücreleri boş kalmalıdır.
LibreOffice yoksa testler atlanır. Dosya Microsoft Excel'de açılarak denenmemiştir.
"""
import os
import re
import shutil
import sys
import tempfile
import unittest
from dataclasses import replace
from datetime import date

import openpyxl

KLASOR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEPO = os.path.dirname(KLASOR)
sys.path.insert(0, KLASOR)
sys.path.insert(0, os.path.join(DEPO, "ortak"))

import demuraj as d  # noqa: E402
import olustur_xlsx as ox  # noqa: E402

try:
    from libreoffice_hesapla import hesapla, soffice_yolu  # noqa: E402
    soffice_yolu()
    LIBREOFFICE_VAR = True
except Exception:  # LibreOfficeYok dahil
    LIBREOFFICE_VAR = False

XLSX = os.path.join(KLASOR, "demuraj-detention.xlsx")
with open(os.path.join(KLASOR, "ornek.json"), encoding="utf-8") as f:
    import json
    ORNEK = d.girdi_oku(json.load(f))

H = ox.HUCRELER
BIRLESIK = [(1, 10, 0), (11, 20, 45), (21, None, 90)]


def girdiyi_yaz(ws, g):
    ws[H["konteyner_sayisi"]] = g.konteyner_sayisi
    ws[H["konteyner_tipi"]] = g.konteyner_tipi
    ws[H["para_birimi"]] = g.para_birimi
    ws[H["mod"]] = ox.MOD_METNI[g.mod] if g.mod in ox.MOD_METNI else g.mod
    ws[H["ilk_gun"]] = "Evet" if g.ilk_gun_sayilir else "Hayır"
    ws[H["kur"]] = g.kur_try
    ws[H["tahliye"]] = g.tahliye
    ws[H["cikis"]] = g.cikis
    ws[H["bos_iade"]] = g.bos_iade
    for tablo in d.TABLOLAR:
        veri = d._satirlari_duzenle(getattr(g, tablo))
        for i, satir in enumerate(ox.satirlar(tablo)):
            bas, bit, ucret = veri[i] if i < len(veri) else (None, None, None)
            ws[f"B{satir}"] = bas
            ws[f"C{satir}"] = bit
            ws[f"D{satir}"] = ucret


def oku_hucreler():
    adresler = [H[k] for k in ("uyari", "demuraj_gun", "detention_gun", "birlesik_gun", "ardiye_gun",
                               "toplam", "toplam_try", "cumle")]
    for tablo in d.TABLOLAR:
        adresler.append(ox.ara_toplam_hucresi(tablo))
        for satir in ox.satirlar(tablo):
            adresler += [f"E{satir}", f"F{satir}"]
    return adresler


def hesapla_senaryo(g):
    gecici = tempfile.mkdtemp(prefix="demuraj-test-")
    try:
        yol = os.path.join(gecici, "senaryo.xlsx")
        shutil.copy(XLSX, yol)
        wb = openpyxl.load_workbook(yol)
        girdiyi_yaz(wb["Hesapla"], g)
        wb.save(yol)
        sayfa = hesapla(yol)["Hesapla"]
        sonuc = {}
        for adres in oku_hucreler():
            m = re.match(r"([A-Z]+)(\d+)", adres)
            sutun = ord(m.group(1)) - 65
            satir = int(m.group(2)) - 1
            try:
                sonuc[adres] = sayfa[satir][sutun]
            except IndexError:
                sonuc[adres] = ""
        return sonuc
    finally:
        shutil.rmtree(gecici, ignore_errors=True)


def sayi(metin):
    return None if metin == "" else float(metin)


def cumleyi_duzelt(metin):
    """Yerel ayar virgül kullanırsa 1234,50 gibi sayıları 1234.50 yapar. Cümledeki diğer virgüllere dokunmaz."""
    return re.sub(r"(\d),(\d{2})(?=\D|$)", r"\1.\2", metin)


def senaryolar():
    s = {}
    s["ornek"] = ORNEK
    s["ilk_gun_hayir"] = replace(ORNEK, ilk_gun_sayilir=False)
    s["birlesik_ardiyeli"] = replace(ORNEK, mod=d.MOD_BIRLESIK, birlesik=BIRLESIK)
    s["birlesik_ardiyesiz_cikissiz"] = replace(ORNEK, mod=d.MOD_BIRLESIK, birlesik=BIRLESIK, ardiye=[], cikis=None)
    s["serbest_sure_icinde"] = replace(ORNEK, cikis=date(2026, 9, 5), bos_iade=date(2026, 9, 8))
    s["kademe_siniri_14"] = replace(ORNEK, ilk_gun_sayilir=False, cikis=date(2026, 9, 15), bos_iade=date(2026, 9, 15),
                                    ardiye=[])
    s["kademe_siniri_15"] = replace(ORNEK, ilk_gun_sayilir=False, cikis=date(2026, 9, 16), bos_iade=date(2026, 9, 16),
                                    ardiye=[])
    s["acik_uclu_son_kademe"] = replace(ORNEK, cikis=date(2026, 10, 10), bos_iade=date(2026, 11, 20))
    s["bes_konteyner_avro_kurlu"] = replace(ORNEK, konteyner_sayisi=5, para_birimi="EUR", kur_try=38.5,
                                            konteyner_tipi="20' DC")
    s["try_para_birimi_kur_yok_sayilir"] = replace(ORNEK, para_birimi="TRY", kur_try=99.0, konteyner_sayisi=1)
    s["bosluklu_bes_satir"] = replace(
        ORNEK, demuraj=[(1, 3, 0), (5, 9, 10), (10, 12, 20), (13, 14, 30), (15, None, 60)], ardiye=[])
    s["ayni_gun_ilk_gun_evet"] = replace(ORNEK, cikis=ORNEK.tahliye, bos_iade=ORNEK.tahliye)
    s["ayni_gun_ilk_gun_hayir"] = replace(ORNEK, cikis=ORNEK.tahliye, bos_iade=ORNEK.tahliye, ilk_gun_sayilir=False)
    return s


def hatali_senaryolar():
    return {
        "cikis_tahliyeden_once": replace(ORNEK, cikis=date(2026, 8, 31)),
        "iade_cikistan_once": replace(ORNEK, bos_iade=date(2026, 9, 10)),
        "iade_eksik": replace(ORNEK, bos_iade=None),
        "tahliye_eksik": replace(ORNEK, tahliye=None),
        "birlesik_ardiye_icin_cikis_eksik": replace(ORNEK, mod=d.MOD_BIRLESIK, birlesik=BIRLESIK, cikis=None),
        "konteyner_sayisi_sifir": replace(ORNEK, konteyner_sayisi=0),
        "kur_eksi": replace(ORNEK, kur_try=-2.0),
        "tarife_bitis_baslangictan_kucuk": replace(ORNEK, detention=[(8, 5, 40)]),
        "tarife_ucreti_bos": replace(ORNEK, demuraj=[(1, 7, None)]),
    }


@unittest.skipUnless(LIBREOFFICE_VAR, "LibreOffice bulunamadı, Excel formülleri LibreOffice ile denenemedi")
class FormulPythonKarsilastirmaTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ox.olustur(XLSX)

    def test_dogru_senaryolar(self):
        for ad, g in senaryolar().items():
            with self.subTest(senaryo=ad):
                s = d.hesapla(g)
                r = hesapla_senaryo(g)
                self.assertEqual(r[H["uyari"]], d.UYARI_YOK)
                for anahtar, deger in (("demuraj_gun", s.demuraj_gun), ("detention_gun", s.detention_gun),
                                       ("birlesik_gun", s.birlesik_gun), ("ardiye_gun", s.ardiye_gun)):
                    self.assertEqual(sayi(r[H[anahtar]]), None if deger is None else float(deger), anahtar)
                self.assertAlmostEqual(float(r[H["toplam"]]), s.toplam, places=6)
                if s.toplam_try is None:
                    self.assertEqual(r[H["toplam_try"]], "")
                else:
                    self.assertAlmostEqual(float(r[H["toplam_try"]]), s.toplam_try, places=6)
                for tablo in d.TABLOLAR:
                    ara = r[ox.ara_toplam_hucresi(tablo)]
                    beklenen_ara = s.ara_toplam[tablo]
                    if beklenen_ara is None:
                        self.assertEqual(ara, "", tablo)
                    else:
                        self.assertAlmostEqual(float(ara), beklenen_ara, places=6)
                    for i, satir in enumerate(ox.satirlar(tablo)):
                        beklenen = s.satirlar[tablo][i]
                        gun, tutar = r[f"E{satir}"], r[f"F{satir}"]
                        if beklenen is None:
                            self.assertEqual((gun, tutar), ("", ""), f"{tablo} satır {i + 1}")
                        else:
                            self.assertEqual(float(gun), beklenen[0], f"{tablo} satır {i + 1} gün")
                            self.assertAlmostEqual(float(tutar), beklenen[1], places=6)
                self.assertEqual(cumleyi_duzelt(r[H["cumle"]]), s.cumle)

    def test_hatali_girdi_uyari_verir_sonuc_gostermez(self):
        for ad, g in hatali_senaryolar().items():
            with self.subTest(senaryo=ad):
                with self.assertRaises(d.GirdiHatasi) as hata:
                    d.hesapla(g)
                r = hesapla_senaryo(g)
                self.assertEqual(r[H["uyari"]], str(hata.exception))
                for anahtar in ("demuraj_gun", "detention_gun", "birlesik_gun", "ardiye_gun", "toplam",
                                "toplam_try", "cumle"):
                    self.assertEqual(r[H[anahtar]], "", anahtar)

    def test_ornek_dosya_ilk_acilista_python_ile_ayni_sonucu_gosterir(self):
        sayfa = hesapla(XLSX)["Hesapla"]
        self.assertEqual(sayfa[26][1], "948")  # B27
        self.assertEqual(cumleyi_duzelt(sayfa[28][1]), d.hesapla(ORNEK).cumle)  # B29


@unittest.skipUnless(LIBREOFFICE_VAR, "LibreOffice bulunamadı")
class YapiTesti(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        ox.olustur(XLSX)

    def test_sayfalar_ve_hakkinda_metni(self):
        wb = openpyxl.load_workbook(XLSX)
        self.assertEqual(wb.sheetnames, ["Hesapla", "Hakkında"])
        metin = "\n".join(str(h.value) for satir in wb["Hakkında"].iter_rows() for h in satir if h.value)
        self.assertIn("2026-10-19", metin)
        self.assertIn("MIT lisansı", metin)
        self.assertIn("LogisticsIQ hazırladı: https://logisticsiq.org", metin)
        self.assertIn("tahmin", metin)

    def test_ayni_girdiden_ayni_dosya_cikar(self):
        gecici = tempfile.mkdtemp(prefix="demuraj-uretim-")
        try:
            a = ox.olustur(os.path.join(gecici, "a.xlsx"))
            b = ox.olustur(os.path.join(gecici, "b.xlsx"))
            with open(a, "rb") as fa, open(b, "rb") as fb:
                self.assertEqual(fa.read(), fb.read())
        finally:
            shutil.rmtree(gecici, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
